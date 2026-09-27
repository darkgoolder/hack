from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.data.comment_parser import CommentParser
from app.data.schema import ImageRecord

DICOM_EXTENSIONS = {".dcm", ".dicom"}


class ExcelLabelError(ValueError):
    pass


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("ё", "е").strip().lower().split())


def _normalize_number(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_binary(value: Any, *, context: str) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        numeric = int(float(str(value).strip().replace(",", ".")))
    except ValueError as exc:
        raise ExcelLabelError(f"Non-binary expert label at {context}: {value!r}") from exc
    if numeric not in (0, 1):
        raise ExcelLabelError(f"Expected 0/1 or blank at {context}, got {value!r}")
    return numeric


def _top_headers(ws) -> list[str]:
    """Return row-1 group headers, expanded through merged ranges.

    Example: B1:D1='Позвоночник' becomes Spine for B,C,D.
    Non-merged cells keep their own values. This also handles sheets with an
    auxiliary column (for example a study UID) between the folder number and
    the clinical groups.
    """
    max_col = ws.max_column
    headers = [ws.cell(1, c).value for c in range(1, max_col + 1)]
    for merged in ws.merged_cells.ranges:
        if merged.min_row == 1 and merged.max_row == 1:
            value = ws.cell(1, merged.min_col).value
            for c in range(merged.min_col, merged.max_col + 1):
                headers[c - 1] = value
    return [_norm(x) for x in headers]


def _find_leaf_columns(ws) -> dict[str, Any]:
    """Discover relevant columns by row-2 leaf headers and row-1 parent groups."""
    max_col = ws.max_column
    parent = _top_headers(ws)
    leaf = [_norm(ws.cell(2, c).value) for c in range(1, max_col + 1)]

    def first_where(predicate):
        for i, value in enumerate(leaf):
            if predicate(value, parent[i]):
                return i
        return None

    folder_col = first_where(lambda l, p: l == "№")
    if folder_col is None:
        folder_col = 1

    comment_col = first_where(lambda l, p: "комментар" in l)

    spine_parent = lambda p: "позвоноч" in p
    right_parent = lambda p: "прав" in p and "бед" in p
    left_parent = lambda p: "лев" in p and "бед" in p
    total_parent = lambda p: "итог" in p

    def col_by_leaf_and_parent(leaf_text: str, parent_predicate):
        for i, value in enumerate(leaf):
            if value == leaf_text and parent_predicate(parent[i]):
                return i
        return None

    cols = {
        "folder": folder_col,
        "comment": comment_col,
        "spine_layout": col_by_leaf_and_parent("корректная укладка", spine_parent),
        "spine_axis": col_by_leaf_and_parent("правильно выравнена ось позвоночника (до 5)", spine_parent),
        "spine_artifact": col_by_leaf_and_parent("наличие посторонних предметов, выраженных артефактов или наложений", spine_parent),
        "right_position_rotation": col_by_leaf_and_parent("позиционирование/ротация", right_parent),
        "right_roi": col_by_leaf_and_parent("корректности области интересов", right_parent),
        "left_position_rotation": col_by_leaf_and_parent("позиционирование/ротация", left_parent),
        "left_roi": col_by_leaf_and_parent("корректности области интересов", left_parent),
        "spine_total": col_by_leaf_and_parent("позвоночник", total_parent),
        "right_total": col_by_leaf_and_parent("проксимальный отдел правого бедра", total_parent),
        "left_total": col_by_leaf_and_parent("проксимальный отдел левого бедра", total_parent),
    }

    missing = [
        name for name, idx in cols.items()
        if name not in {"comment"} and idx is None
    ]
    if missing:
        raise ExcelLabelError(
            "Could not locate required XLSX columns: " + ", ".join(missing) +
            ". The parser expects the two-level header described in the project specification."
        )
    return cols


def parse_excel(excel_path: Path, sheet_name: str | None = None) -> dict[str, dict[str, Any]]:
    wb = load_workbook(excel_path, data_only=True, read_only=False)
    ws = wb[sheet_name] if sheet_name else wb.active
    if ws.max_row < 3:
        raise ExcelLabelError("Excel must contain two header rows and at least one data row")

    cols = _find_leaf_columns(ws)
    result: dict[str, dict[str, Any]] = {}

    for excel_row_no in range(3, ws.max_row + 1):
        folder_value = ws.cell(excel_row_no, cols["folder"] + 1).value
        folder_id = _normalize_number(folder_value)
        if not folder_id:
            continue

        def cell(col_name: str):
            idx = cols[col_name]
            return ws.cell(excel_row_no, idx + 1).value

        comment = None
        if cols["comment"] is not None:
            raw = cell("comment")
            comment = str(raw).strip() if raw is not None and str(raw).strip() else None

        record: dict[str, Any] = {
            "excel_row": excel_row_no,
            "folder_id": folder_id,
            "comment": comment,
            "spine": {
                "layout": _to_binary(cell("spine_layout"), context=f"row {excel_row_no}, spine_layout"),
                "axis": _to_binary(cell("spine_axis"), context=f"row {excel_row_no}, spine_axis"),
                "artifact": _to_binary(cell("spine_artifact"), context=f"row {excel_row_no}, spine_artifact"),
                "total": _to_binary(cell("spine_total"), context=f"row {excel_row_no}, spine_total"),
            },
            "right_hip": {
                "position_rotation": _to_binary(cell("right_position_rotation"), context=f"row {excel_row_no}, right_position_rotation"),
                "roi": _to_binary(cell("right_roi"), context=f"row {excel_row_no}, right_roi"),
                "total": _to_binary(cell("right_total"), context=f"row {excel_row_no}, right_total"),
            },
            "left_hip": {
                "position_rotation": _to_binary(cell("left_position_rotation"), context=f"row {excel_row_no}, left_position_rotation"),
                "roi": _to_binary(cell("left_roi"), context=f"row {excel_row_no}, left_roi"),
                "total": _to_binary(cell("left_total"), context=f"row {excel_row_no}, left_total"),
            },
        }

        warnings: list[str] = []
        for region in ("spine", "right_hip", "left_hip"):
            details = [v for k, v in record[region].items() if k != "total" and v is not None]
            total = record[region]["total"]
            if details and total is not None and max(details) != total:
                warnings.append(
                    f"{region}: total={total} differs from max(detail_labels)={max(details)}"
                )
        record["validation_warnings"] = warnings
        result[folder_id] = record

    return result


def classify_file(file_name: str) -> tuple[str, str | None]:
    normalized = file_name.strip().lower()
    mapping = {
        "spine.dcm": ("spine", None),
        "left_hip.dcm": ("left_hip", "L"),
        "right_hip.dcm": ("right_hip", "R"),
    }
    if normalized in mapping:
        return mapping[normalized]
    raise ValueError(
        f"Unexpected DICOM filename: {file_name!r}. "
        "Expected spine.dcm, left_hip.dcm or right_hip.dcm."
    )


def build_records(dataset_root: Path, labels: dict[str, dict[str, Any]], comment_parser: CommentParser) -> tuple[list[ImageRecord], list[str]]:
    # Local import: the Excel parser and its unit tests do not require pydicom,
    # but actual DICOM auditing does.
    from app.dicom.loader import get_dicom_info

    records: list[ImageRecord] = []
    warnings: list[str] = []
    folders = sorted((p for p in dataset_root.iterdir() if p.is_dir()), key=lambda p: p.name)
    seen_folders: set[str] = set()

    for folder in folders:
        folder_id = folder.name
        seen_folders.add(folder_id)
        if folder_id not in labels:
            warnings.append(f"Folder {folder_id}: no corresponding Excel row")
            continue

        row = labels[folder_id]
        comment = row.get("comment")
        auxiliary_tags = comment_parser.parse(comment)
        if comment_parser.is_ambiguous(comment):
            warnings.append(f"Folder {folder_id}: ambiguous comment {comment!r} stored as auxiliary-only")

        files = sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in DICOM_EXTENSIONS)
        expected_names = {"spine.dcm", "left_hip.dcm", "right_hip.dcm"}
        actual_names = {p.name.lower() for p in files}
        missing = expected_names - actual_names
        if missing:
            warnings.append(f"Folder {folder_id}: missing expected files: {sorted(missing)}")

        for path in files:
            try:
                anatomy, laterality = classify_file(path.name)
            except ValueError as exc:
                warnings.append(f"Folder {folder_id}: {exc}")
                continue

            try:
                info = get_dicom_info(path)
                readable = True
                error = None
            except Exception as exc:  # noqa: BLE001
                info = None
                readable = False
                error = f"{type(exc).__name__}: {exc}"
                warnings.append(f"{folder_id}/{path.name}: DICOM read failed: {error}")

            expert = row[anatomy]
            if anatomy == "spine":
                violation_names = (
                    ("SPINE_LAYOUT", expert["layout"]),
                    ("SPINE_AXIS_DEVIATION", expert["axis"]),
                    ("SPINE_ARTIFACT", expert["artifact"]),
                )
                record = ImageRecord(
                    folder_id=folder_id,
                    image_path=path.relative_to(dataset_root),
                    file_name=path.name,
                    anatomy=anatomy,
                    laterality=laterality,
                    projection=info.projection if info else None,
                    study_uid=info.study_uid if info else None,
                    series_uid=info.series_uid if info else None,
                    sop_instance_uid=info.sop_instance_uid if info else None,
                    spine_layout=expert["layout"],
                    spine_axis=expert["axis"],
                    spine_artifact=expert["artifact"],
                    quality_class=expert["total"],
                    violation_type=";".join(name for name, value in violation_names if value == 1) or None,
                    expert_comment=comment,
                    auxiliary_tags=auxiliary_tags,
                    dicom_readable=readable,
                    error=error,
                )
            else:
                violation_names = (
                    ("HIP_POSITION_ROTATION", expert["position_rotation"]),
                    ("HIP_ROI", expert["roi"]),
                )
                record = ImageRecord(
                    folder_id=folder_id,
                    image_path=path.relative_to(dataset_root),
                    file_name=path.name,
                    anatomy=anatomy,
                    laterality=laterality,
                    projection=info.projection if info else None,
                    study_uid=info.study_uid if info else None,
                    series_uid=info.series_uid if info else None,
                    sop_instance_uid=info.sop_instance_uid if info else None,
                    hip_position_rotation=expert["position_rotation"],
                    hip_roi=expert["roi"],
                    quality_class=expert["total"],
                    violation_type=";".join(name for name, value in violation_names if value == 1) or None,
                    expert_comment=comment,
                    auxiliary_tags=auxiliary_tags,
                    dicom_readable=readable,
                    error=error,
                )
            records.append(record)

    for folder_id in labels:
        if folder_id not in seen_folders:
            warnings.append(f"Excel row for folder {folder_id}: folder not found in dataset")

    return records, warnings


CANONICAL_COLUMN_NAMES = {
    "anatomy": "anatomical_region",
    "expert_comment": "study_comments",
}


def write_manifest(records: list[ImageRecord], output_csv: Path) -> None:
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    if not records:
        raise ValueError("No records to write")

    rows = [
        {CANONICAL_COLUMN_NAMES.get(k, k): v for k, v in record.to_dict().items()}
        for record in records
    ]
    fieldnames = list(rows[0].keys())
    with output_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a validated DICOM/Excel manifest for DXA QC")
    parser.add_argument("--dataset-root", type=Path, required=True)
    parser.add_argument("--excel", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("data/manifests/dataset_manifest.csv"))
    parser.add_argument("--warnings-output", type=Path, default=Path("data/manifests/dataset_warnings.txt"))
    parser.add_argument("--sheet", type=str, default=None)
    parser.add_argument("--labels-config", type=Path, default=Path("configs/labels.yaml"))
    args = parser.parse_args()

    labels = parse_excel(args.excel, sheet_name=args.sheet)
    comment_parser = CommentParser(args.labels_config)
    records, warnings = build_records(args.dataset_root, labels, comment_parser)
    write_manifest(records, args.output)
    args.warnings_output.parent.mkdir(parents=True, exist_ok=True)
    args.warnings_output.write_text("\n".join(warnings), encoding="utf-8")

    print(f"Manifest written: {args.output}")
    print(f"Records: {len(records)}")
    print(f"Warnings: {len(warnings)}")
    if warnings:
        print(f"Warnings written: {args.warnings_output}")


if __name__ == "__main__":
    main()
