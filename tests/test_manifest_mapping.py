from pathlib import Path

import openpyxl

from app.data.build_manifest import parse_excel


def test_parse_expected_two_row_header(tmp_path: Path) -> None:
    xlsx = tmp_path / "labels.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.append([
        "№", "Позвоночник", None, None, "Проксимальный отдел правого бедра", None,
        "Проксимальный отдел левого бедра", None, "Итог", None, None, "Комментарий"
    ])
    ws.merge_cells("B1:D1")
    ws.merge_cells("E1:F1")
    ws.merge_cells("G1:H1")
    ws.merge_cells("I1:K1")
    ws.append([
        "№", "корректная укладка", "правильно выравнена ось позвоночника (до 5)",
        "наличие посторонних предметов, выраженных артефактов или наложений",
        "позиционирование/ротация", "корректности области интересов",
        "позиционирование/ротация", "корректности области интересов",
        "Позвоночник", "Проксимальный отдел правого бедра",
        "Проксимальный отдел левого бедра", "Комментарий"
    ])
    ws.append([1, 0, 1, 0, 1, 0, 0, 1, 1, 1, 1, "сколиоз"])
    wb.save(xlsx)

    result = parse_excel(xlsx)
    row = result["1"]
    assert row["spine"] == {"layout": 0, "axis": 1, "artifact": 0, "total": 1}
    assert row["right_hip"] == {"position_rotation": 1, "roi": 0, "total": 1}
    assert row["left_hip"] == {"position_rotation": 0, "roi": 1, "total": 1}
    assert row["comment"] == "сколиоз"
